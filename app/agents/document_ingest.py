"""Convert uploaded documents into clean plain text.

Strategy (cheapest first):
  - PDF              -> pypdf text extraction; if the PDF is scanned or yields
                        almost no text, render the first pages to images and read
                        them with the Groq vision model.
  - Images           -> Groq vision (portfolio screenshots, certificates).
  - DOCX             -> python-docx paragraphs and tables.
  - Text/code/etc.   -> read directly (indentation preserved for code).

The profile text model (GPT-OSS 120B) is text-only, so images and scanned PDFs
are routed to a Groq vision model instead.
"""

import base64
import logging
from pathlib import Path

import pymupdf
from docx import Document
from langchain_core.messages import HumanMessage
from pypdf import PdfReader

from app.agents.llm import get_vision_model, message_text
from app.agents.text_cleaning import clean_text

logger = logging.getLogger(__name__)

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tiff"}

CODE_EXTENSIONS = {
    ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".c", ".cpp", ".h", ".hpp",
    ".cs", ".go", ".rb", ".rs", ".php", ".sql", ".sh", ".kt", ".swift", ".scala",
    ".r", ".m", ".pl", ".lua", ".dart", ".vue", ".svelte",
}

# Below this many characters we assume a PDF has no extractable text layer.
MIN_PDF_TEXT_CHARS = 200

# Groq vision accepts at most 3 images per request, so cap scanned-PDF pages.
MAX_VISION_PAGES = 3

_FILE_READ_PROMPT = (
    "Extract all text and meaningful content from the provided image(s). Preserve "
    "names, dates, skills, job titles, education, and any structured details. "
    "Return plain text only, with no commentary."
)


def _image_block(media_type: str, raw: bytes) -> dict:
    """Build an OpenAI-style base64 image block for the Groq vision API."""
    encoded = base64.b64encode(raw).decode("ascii")
    return {
        "type": "image_url",
        "image_url": {"url": f"data:{media_type};base64,{encoded}"},
    }


def _read_images_with_vision(images: list[tuple[str, bytes]]) -> str:
    """Ask the Groq vision model to read up to 3 images and return their text."""
    if not images:
        return ""

    content: list[dict] = [{"type": "text", "text": _FILE_READ_PROMPT}]
    for media_type, raw in images[:MAX_VISION_PAGES]:
        content.append(_image_block(media_type, raw))

    try:
        response = get_vision_model().invoke([HumanMessage(content=content)])
        return message_text(response)
    except Exception as error:
        # If the vision call fails (e.g. a transient rate limit), log it and fall
        # back to whatever text we already have.
        logger.warning("Groq vision read failed for %d image(s): %s", len(content) - 1, error)
        return ""


def _render_pdf_pages(path: Path, max_pages: int = MAX_VISION_PAGES) -> list[tuple[str, bytes]]:
    """Render the first pages of a PDF to PNG images for the vision model."""
    images: list[tuple[str, bytes]] = []
    try:
        with pymupdf.open(str(path)) as document:
            for page_index in range(min(max_pages, document.page_count)):
                page = document.load_page(page_index)
                pixmap = page.get_pixmap(dpi=150)
                images.append(("image/png", pixmap.tobytes("png")))
    except Exception:
        return []
    return images


def _extract_pdf(path: Path) -> str:
    text = ""
    try:
        reader = PdfReader(str(path))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception:
        text = ""

    if len(text.strip()) >= MIN_PDF_TEXT_CHARS:
        return clean_text(text)

    # Scanned / image-only PDF: render pages and let the vision model read them.
    vision_text = _read_images_with_vision(_render_pdf_pages(path))
    return clean_text(vision_text or text)


def _extract_docx(path: Path) -> str:
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return clean_text("\n".join(parts))


def _extract_image(path: Path, mime_type: str | None) -> str:
    media_type = mime_type or f"image/{path.suffix.lstrip('.').lower()}"
    return clean_text(_read_images_with_vision([(media_type, path.read_bytes())]))


def _extract_plain_text(path: Path, preserve_indentation: bool) -> str:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    return clean_text(raw, preserve_indentation=preserve_indentation)


def extract_document_text(file_path: str, mime_type: str | None, file_name: str) -> str:
    """Return cleaned plain text for a single uploaded document."""
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf" or mime_type == "application/pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix in IMAGE_EXTENSIONS or (mime_type or "").startswith("image/"):
        return _extract_image(path, mime_type)

    return _extract_plain_text(path, preserve_indentation=suffix in CODE_EXTENSIONS)
